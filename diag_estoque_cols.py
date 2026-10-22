"""
Diagnostico completo das colunas do Estoque_Geral
Le sem data_only=True para ver valores reais E formulas
"""
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

XLSX = r"C:\Users\valqu\AppData\Local\Temp\Planilha_Mestre_temp.xlsx"

print("=== LEITURA SEM data_only (valores hardcoded + fórmulas como string) ===")
wb1 = openpyxl.load_workbook(XLSX, data_only=False)
ws1 = wb1['Estoque_Geral']

print(f"\nTotal de colunas: {ws1.max_column}, Total de linhas: {ws1.max_row}")
print("\n--- CABEÇALHOS (linha 1) ---")
for col in range(1, ws1.max_column + 1):
    cell = ws1.cell(row=1, column=col)
    print(f"  Col {cell.column_letter} (idx={col}): {repr(cell.value)}")

print("\n--- PRIMEIRAS 3 LINHAS DE DADOS (linhas 2-4) ---")
for row_idx in range(2, 5):
    row = ws1[row_idx]
    print(f"\n  Linha {row_idx}:")
    for cell in row:
        if cell.value is not None:
            print(f"    {cell.column_letter} (idx={cell.column}): {repr(cell.value)}")

print("\n\n=== LEITURA COM data_only=True (valores calculados) ===")
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2['Estoque_Geral']

print("\n--- PRIMEIRAS 3 LINHAS DE DADOS (linhas 2-4) ---")
for row_idx in range(2, 5):
    row = ws2[row_idx]
    print(f"\n  Linha {row_idx}:")
    for cell in row:
        print(f"    {cell.column_letter} (idx={cell.column}): {repr(cell.value)}")

print("\n\n=== COMPARAÇÃO DIRETA: linha 2 ambos modos ===")
print("\nSem data_only:")
for cell in ws1[2]:
    print(f"  {cell.column_letter}: {repr(cell.value)}")
print("\nCom data_only:")
for cell in ws2[2]:
    print(f"  {cell.column_letter}: {repr(cell.value)}")
