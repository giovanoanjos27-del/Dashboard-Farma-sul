"""
explore_excel_structure.py
Explora a estrutura da 1-Tabela-1 para entender colunas e períodos disponíveis.
"""
import openpyxl, os, re
from collections import defaultdict

BASE = r"C:\Users\valqu\OneDrive\Documentos\Claude\Projects\projeto Dashboard_farma_sul"
XLSX = os.path.join(BASE, "base_temp_copy.xlsx")
SHEET = "1-Tabela-1"

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb[SHEET]

# Ler as primeiras 20 linhas para entender cabeçalhos
print("=== PRIMEIRAS 20 LINHAS ===")
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
    print(f"Row {i+1}: {list(row[:12])}")

# Coletar set de valores únicos em algumas colunas-chave
markets = set()
periods = set()
brands = set()
value_types = set()

print("\n=== ESCANEANDO DADOS ===")
for i, row in enumerate(ws.iter_rows(min_row=10, values_only=True)):
    if not row[0]:
        continue
    # Tentar identificar a estrutura
    for j, val in enumerate(row):
        if val and isinstance(val, str):
            # Coluna 0 = market (geralmente)
            if i == 0:
                print(f"Col {j}: '{val}'")
    if i > 5:
        break

# Identificar colunas pelo conteúdo
print("\n=== IDENTIFICANDO COLUNAS (linhas 10-50) ===")
col_samples = defaultdict(set)
for i, row in enumerate(ws.iter_rows(min_row=10, max_row=50, values_only=True)):
    for j, val in enumerate(row):
        if val is not None:
            col_samples[j].add(str(val)[:30])

for j in sorted(col_samples.keys()):
    vals = list(col_samples[j])[:5]
    print(f"Col {j}: {vals}")

wb.close()
