#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_market_total.py
Extrai FYTD25 e FYTD26 do total do mercado Personal Repellent (TODAS as marcas)
por mercado, a partir do Excel 1-Tabela-1.
Salva em fytd_market_total.json
"""
import json, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Instale openpyxl: pip install openpyxl")

EXCEL_PATH = Path(r"C:\Users\valqu\OneDrive\Documentos\Claude\Projects\projeto Dashboard_farma_sul\Dashboard_farma_sul\_Recursos\Giovano SCJ\Base Giovano 26-04.xlsx")
OUT_JSON   = Path(r"C:\Users\valqu\OneDrive\Documentos\Claude\Projects\projeto Dashboard_farma_sul\fytd_market_total.json")
SHEET_NAME = "1-Tabela-1"

print(f"Abrindo {EXCEL_PATH.name} ...")
wb = openpyxl.load_workbook(str(EXCEL_PATH), read_only=True, data_only=True)

if SHEET_NAME not in wb.sheetnames:
    sys.exit(f"Aba '{SHEET_NAME}' nao encontrada. Abas: {wb.sheetnames}")

ws = wb[SHEET_NAME]
print("Lendo cabecalhos...")

# Detect header row (row 9 based on previous session)
header_row = None
col_market = col_category = col_brand = col_period = col_val = col_un = None

for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
    for i, cell in enumerate(row):
        if cell is None: continue
        s = str(cell).strip().upper()
        if s in ("MARKETS", "MARKET"):
            col_market = i
        elif s == "CATEGORY":
            col_category = i
        elif s == "BRAND":
            col_brand = i
        elif s in ("DATA", "DATE", "PERIOD"):
            col_period = i
        elif s == "$":
            col_val = i
        elif s in ("(UNID)", "UNID", "UNITS"):
            col_un = i
    if col_market is not None and col_period is not None and col_val is not None:
        header_row = [str(c).strip() if c else "" for c in row]
        print(f"Header: market={col_market}, cat={col_category}, brand={col_brand}, period={col_period}, val={col_val}, un={col_un}")
        break

if header_row is None:
    sys.exit("Cabecalho nao encontrado")

# totals[market][period] = {val, un}
totals = {}

row_count = 0
skip_count = 0
for row in ws.iter_rows(min_row=ws.min_row + 1, values_only=True):
    row_count += 1
    if row_count % 10000 == 0:
        sys.stdout.write(f"  ...{row_count} linhas processadas\n")
        sys.stdout.flush()

    # Category filter
    if col_category is not None:
        cat = row[col_category]
        if cat is None or str(cat).strip().upper() != "PERSONAL REPELLENT":
            skip_count += 1
            continue

    # Period filter
    period = row[col_period]
    if period is None:
        continue
    p = str(period).strip().upper()
    if p not in ("FYTD25", "FYTD26"):
        continue

    mkt = row[col_market]
    if mkt is None:
        continue
    mkt = str(mkt).strip()
    if not mkt:
        continue

    # Value (all brands — no brand filter)
    val_raw = row[col_val]
    un_raw  = row[col_un] if col_un is not None else None

    try:
        val = float(val_raw) / 1e6 if val_raw is not None else 0.0
    except (TypeError, ValueError):
        val = 0.0
    try:
        un = float(un_raw) / 1e3 if un_raw is not None else 0.0
    except (TypeError, ValueError):
        un = 0.0

    if mkt not in totals:
        totals[mkt] = {}
    if p not in totals[mkt]:
        totals[mkt][p] = {"val": 0.0, "un": 0.0}

    totals[mkt][p]["val"] += val
    totals[mkt][p]["un"]  += un

wb.close()
print(f"Total de linhas lidas: {row_count} | Ignoradas (cat): {skip_count}")
print(f"Mercados encontrados: {len(totals)}")

# Round values
for mkt in totals:
    for p in totals[mkt]:
        totals[mkt][p]["val"] = round(totals[mkt][p]["val"], 6)
        totals[mkt][p]["un"]  = round(totals[mkt][p]["un"],  3)

# Print sample
sample_mkt = "T.BR Farma"
if sample_mkt in totals:
    f25 = totals[sample_mkt].get("FYTD25", {})
    f26 = totals[sample_mkt].get("FYTD26", {})
    cres = (f26.get("val",0)/f25.get("val",1) - 1)*100 if f25.get("val",0)>0 else 0
    sys.stdout.write(f"{sample_mkt}: F25={f25.get('val',0):.3f}MM  F26={f26.get('val',0):.3f}MM  Cres={cres:+.2f}%%\n")
    sys.stdout.flush()

with open(str(OUT_JSON), "w", encoding="utf-8") as f:
    json.dump(totals, f, indent=2, ensure_ascii=False)
print(f"JSON salvo: {OUT_JSON.name}")
print("DONE!")
